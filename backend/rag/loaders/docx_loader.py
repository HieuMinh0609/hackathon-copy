"""Đọc file Word (.docx) chuẩn UTF-8 bằng python-docx, hỗ trợ đa ngôn ngữ (VI/EN/KR)."""
from __future__ import annotations

from pathlib import Path
from langchain_core.documents import Document


def load_docx(path: Path) -> list[Document]:
    """Trích xuất nội dung từ file docx bao gồm các đoạn văn bản (paragraphs) và bảng biểu (tables)."""
    import docx

    try:
        doc = docx.Document(str(path))
    except Exception as exc:
        raise RuntimeError(f"Không thể đọc file Word '{path.name}': {exc}") from exc

    elements: list[str] = []

    # 1. Trích xuất các đoạn văn bản
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            elements.append(txt)

    # 2. Trích xuất bảng biểu
    for table_idx, table in enumerate(doc.tables, start=1):
        table_lines = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(row_cells):
                table_lines.append(" | ".join(row_cells))
        if table_lines:
            elements.append(f"[Bảng {table_idx}]\n" + "\n".join(table_lines))

    full_text = "\n\n".join(elements).strip()
    if not full_text:
        return []

    return [
        Document(
            page_content=full_text,
            metadata={"source": path.name, "file_type": "docx"},
        )
    ]
