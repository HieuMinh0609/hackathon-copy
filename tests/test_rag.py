from pathlib import Path
import pytest

pytest.importorskip("langchain_text_splitters")

from langchain_core.documents import Document
from langchain_core.embeddings import FakeEmbeddings
from langchain_chroma import Chroma

from backend.rag.chunker import split_documents
from backend.rag.ingest import discover_files, infer_role, ingest_files, load_file
from backend.rag.retriever import format_multilingual_context, retrieve


def test_infer_role(tmp_path: Path):
    (tmp_path / "hr").mkdir()
    assert infer_role(tmp_path / "hr" / "a.md", tmp_path) == "hr"
    assert infer_role(tmp_path / "root.md", tmp_path) == "onboarding"


def test_discover_and_load_markdown(tmp_path: Path):
    f = tmp_path / "onboarding" / "a.md"
    f.parent.mkdir()
    f.write_text("# Tiêu đề Tiếng Việt\n한국어 내용 온보딩", encoding="utf-8")
    (tmp_path / "ignore.exe").write_text("x")
    assert discover_files(tmp_path) == [f]
    docs = load_file(f)
    assert len(docs) == 1
    assert docs[0].metadata["source"] == "a.md"
    assert "한국어" in docs[0].page_content


def test_load_docx(tmp_path: Path):
    import docx
    doc_path = tmp_path / "guide.docx"
    doc = docx.Document()
    doc.add_paragraph("Tài liệu hướng dẫn Tiếng Việt")
    doc.add_paragraph("프로젝트 가이드라인 한국어 (Hangul Korean)")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Tên mục"
    table.rows[0].cells[1].text = "Giá trị"
    doc.save(str(doc_path))

    loaded = load_file(doc_path)
    assert len(loaded) == 1
    assert loaded[0].metadata["source"] == "guide.docx"
    assert loaded[0].metadata["file_type"] == "docx"
    assert "Tiếng Việt" in loaded[0].page_content
    assert "한국어" in loaded[0].page_content
    assert "Hangul" in loaded[0].page_content


def test_chunker_splits_long_text(settings, tmp_path: Path):
    f = tmp_path / "long.md"
    f.write_text("\n\n".join(f"Đoạn {i}. " + "nội dung " * 60 for i in range(10)), encoding="utf-8")
    chunks = split_documents(settings, load_file(f))
    assert len(chunks) > 1
    assert all(len(c.page_content) <= settings.chunk_size + 200 for c in chunks)


def test_format_multilingual_context():
    doc1 = Document(
        page_content="Nội dung Tiếng Việt quy định dự án",
        metadata={"source": "quy_dinh.pdf", "page": 1, "file_type": "pdf", "role": "onboarding"},
    )
    doc2 = Document(
        page_content="한국어 프로젝트 온보딩 체크리스트",
        metadata={"source": "onboarding_kr.docx", "file_type": "docx", "role": "backend_developer"},
    )
    formatted = format_multilingual_context([doc1, doc2])
    assert "quy_dinh.pdf" in formatted
    assert "Trang 1" in formatted
    assert "onboarding_kr.docx" in formatted
    assert "한국어" in formatted


def test_multithreaded_ingest_with_fake_embeddings(settings, tmp_path: Path, monkeypatch):
    # Dùng FakeEmbeddings 1024 chiều (giống bge-m3) để test đa luồng và batching offline
    fake_emb = FakeEmbeddings(size=1024)
    fake_store = Chroma(
        collection_name="test_col",
        embedding_function=fake_emb,
        persist_directory=str(settings.chroma_dir),
    )
    monkeypatch.setattr("backend.rag.vector_store.get_vector_store", lambda s: fake_store)

    f1 = tmp_path / "onboarding" / "doc1.md"
    f1.parent.mkdir(parents=True, exist_ok=True)
    f1.write_text("\n\n".join(f"Câu {i}: Nội dung đào tạo dự án." for i in range(20)), encoding="utf-8")

    f2 = tmp_path / "onboarding" / "doc2.md"
    f2.write_text("\n\n".join(f"Câu {i}: 한국어 프로젝트 안내." for i in range(20)), encoding="utf-8")

    progress_calls = []

    def progress(i: int, total: int, msg: str):
        progress_calls.append((i, total, msg))

    report = ingest_files(
        settings,
        [f1, f2],
        root=tmp_path,
        reset=False,
        progress=progress,
        batch_size=16,
        max_workers=2,
    )

    assert report.files == 2
    assert report.chunks > 0
    assert len(report.errors) == 0
    assert len(progress_calls) > 0
    assert fake_store._collection.count() == report.chunks


def test_ingest_and_retrieve_by_project(settings, tmp_path: Path, monkeypatch):
    from backend.services.knowledge_service import KnowledgeService

    fake_emb = FakeEmbeddings(size=1024)
    fake_store = Chroma(
        collection_name="test_proj_col",
        embedding_function=fake_emb,
        persist_directory=str(settings.chroma_dir),
    )
    monkeypatch.setattr("backend.rag.vector_store.get_vector_store", lambda s: fake_store)

    service = KnowledgeService(settings)

    # 1. Tạo các file tài liệu cho 2 dự án khác nhau
    p1_file = tmp_path / "prj1_db.md"
    p1_file.write_text("Hướng dẫn kết nối PostgreSQL dự án 1.", encoding="utf-8")

    p2_file = tmp_path / "prj2_hr.md"
    p2_file.write_text("Chính sách làm việc từ xa dự án 2.", encoding="utf-8")

    # 2. Nạp dữ liệu theo project_code
    rep1 = service.ingest_project("PRJ_ONE", [p1_file])
    rep2 = service.ingest_project("PRJ_TWO", [p2_file])
    assert rep1.chunks > 0
    assert rep2.chunks > 0

    # 3. Liệt kê danh sách dự án
    projects = service.list_projects()
    assert "PRJ_ONE" in projects
    assert "PRJ_TWO" in projects

    # 4. Lấy trực tiếp tài liệu theo mã dự án
    docs = service.get_project_docs("PRJ_ONE")
    assert len(docs) == rep1.chunks
    assert "PostgreSQL" in docs[0].page_content
    assert docs[0].metadata["project_code"] == "PRJ_ONE"

    # 5. Tìm kiếm ngữ nghĩa trong phạm vi dự án
    results = service.search_project("PRJ_ONE", "PostgreSQL", k=2)
    assert len(results) > 0
    assert results[0].metadata["project_code"] == "PRJ_ONE"

    # 6. Xóa theo mã dự án
    deleted = service.delete_project_docs("PRJ_TWO")
    assert deleted > 0
    assert service.get_project_docs("PRJ_TWO") == []
    assert "PRJ_TWO" not in service.list_projects()


